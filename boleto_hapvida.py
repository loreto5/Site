import os
import requests
import logging
import time
import re
import pandas as pd
from bs4 import BeautifulSoup
from urllib.parse import urljoin
from playwright.sync_api import sync_playwright
import certifi
from datetime import datetime
from concurrent.futures import ThreadPoolExecutor, as_completed
import openpyxl
from openpyxl.styles import PatternFill, Font
import queue
from contextlib import contextmanager
import threading

# Configurar logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('boleto_hapvida.log', encoding='utf-8'),
        logging.StreamHandler()
    ]
)
logger = logging.getLogger(__name__)

# Configurações
BASE_URL = "https://webhap.hapvida.com.br/pls/webhap/"
LOGIN_URL = BASE_URL + "webNewBoletoEmpresa.loginValidacao"
FATURA_URL = BASE_URL + "webNewBoletoEmpresa.obrigacaoDados"
CONFIRMACAO_URL = BASE_URL + "webNewBoletoEmpresa.validaObrigacao"

# Coluna onde será escrito o status (0-indexado, então coluna C = 2)
STATUS_COLUMN = 2

# Número máximo de tentativas para cada requisição
MAX_RETRIES = 3
# Tempo de espera entre tentativas (em segundos)
RETRY_DELAY = 0
# Número máximo de workers para processamento paralelo
MAX_WORKERS = 15

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/91.0.4472.124 Safari/537.36",
    "Content-Type": "application/x-www-form-urlencoded",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/webp,*/*;q=0.8",
    "Accept-Language": "pt-BR,pt;q=0.9,en-US;q=0.8,en;q=0.7"
}

# Cores para formatação do Excel
COR_SUCESSO = "90EE90"  # Verde claro
COR_ERRO = "FFCCCB"     # Vermelho claro

# Variáveis globais para rastrear progresso
progress_lock = threading.Lock()
global_processed = 0
global_total = 0

# Gerenciador de pool de sessões HTTP
class SessionPool:
    def __init__(self, pool_size=MAX_WORKERS):
        self.pool_size = pool_size
        self.sessions = [requests.Session() for _ in range(pool_size)]
        for session in self.sessions:
            session.headers.update(HEADERS)

    @contextmanager
    def get_session(self):
        session = self.sessions.pop() if self.sessions else requests.Session()
        session.headers.update(HEADERS)
        try:
            yield session
        finally:
            if len(self.sessions) < self.pool_size:
                self.sessions.append(session)
            else:
                session.close()

    def close_all(self):
        for session in self.sessions:
            session.close()
        self.sessions = []

class PlaywrightManager:
    """Classe para gerenciar instâncias do Playwright para reutilização"""
    def __init__(self):
        self.playwright = None
        self.browser = None
        self.initialized = False
    
    def initialize(self):
        """Inicializa o Playwright e o navegador se ainda não estiverem inicializados"""
        if not self.initialized:
            self.playwright = sync_playwright().start()
            self.browser = self.playwright.chromium.launch(headless=True)
            self.initialized = True
    
    def create_context(self):
        """Cria um novo contexto de navegador"""
        self.initialize()
        return self.browser.new_context()
    
    def close(self):
        """Fecha o navegador e o Playwright"""
        if self.initialized:
            if self.browser:
                self.browser.close()
            if self.playwright:
                self.playwright.stop()
            self.initialized = False

def gerar_pdf_com_browser_reutilizavel(context, url, cookies, output_path, retries=MAX_RETRIES):
    """Gera um PDF a partir de uma URL usando um contexto reutilizável"""
    for attempt in range(1, retries + 1):
        try:
            context.clear_cookies()
            context.add_cookies(cookies)
            page = context.new_page()
            page.goto(url, wait_until="networkidle", timeout=60000)
            if page.title() == "Error" or "Erro" in page.content():
                page.close()
                raise Exception("Página de erro detectada")
            page.wait_for_timeout(2000)
            page.pdf(path=output_path, format="A4", print_background=True)
            page.close()
            if os.path.exists(output_path) and os.path.getsize(output_path) > 10000:
                return True
            else:
                raise Exception("PDF gerado com tamanho insuficiente")
        except Exception as e:
            logger.warning(f"Tentativa {attempt} falhou: {str(e)}")
            if attempt < retries:
                time.sleep(RETRY_DELAY)
            else:
                logger.error(f"Falha ao gerar PDF após {retries} tentativas: {str(e)}")
                return False
    return False

def request_with_retry(session, method, url, data=None, timeout=(10, 60), retries=MAX_RETRIES):
    """Faz requisições HTTP com sistema de retentativas"""
    for attempt in range(1, retries + 1):
        try:
            if method.lower() == 'get':
                response = session.get(url, headers=HEADERS, timeout=timeout, verify=certifi.where())
            else:  # post
                response = session.post(url, data=data, headers=HEADERS, timeout=timeout, verify=certifi.where())
            response.raise_for_status()
            return response
        except (requests.exceptions.RequestException, requests.exceptions.HTTPError) as e:
            logger.warning(f"Tentativa {attempt} falhou para {url}: {str(e)}")
            if attempt < retries:
                time.sleep(RETRY_DELAY)
            else:
                logger.error(f"Falha após {retries} tentativas para {url}: {str(e)}")
                raise

def validar_data_vencimento(data_vencimento):
    """Valida o formato da data de vencimento e retorna True se válida"""
    if not re.match(r'\d{2}/\d{2}/\d{4}', data_vencimento):
        return False
    try:
        datetime.strptime(data_vencimento, '%d/%m/%Y')
        return True
    except ValueError:
        return False

def processar_boleto(pw_manager, context, session_pool, empresa, senha, data_vencimento, nome_destino, save_dir):
    """Processa um boleto específico para uma empresa"""
    if not validar_data_vencimento(data_vencimento):
        return f"❌ Data inválida para {empresa}: {data_vencimento}"

    with session_pool.get_session() as sessao:
        os.makedirs(save_dir, exist_ok=True)
        try:
            empresa_code = empresa.upper()[:5] if empresa else ""
            senha_clean = senha[:10] if senha else ""
            if not empresa_code or not senha_clean:
                return f"❌ Credenciais inválidas para {empresa}"
            login_data = {
                "pCodigo": empresa_code,
                "pSenha": senha_clean,
                "pIdSessao": "",
                "pNoCache": "",
                "pTokenCaptcha": ""
            }
            request_with_retry(sessao, 'get', BASE_URL + "webNewBoletoEmpresa.login")
            login_response = request_with_retry(sessao, 'post', LOGIN_URL, data=login_data)
            if "ERRO" in login_response.text.upper() or "FALHA" in login_response.text.upper():
                erro_match = re.search(r'<div[^>]*class="erro"[^>]*>(.*?)</div>', login_response.text, re.DOTALL)
                erro_msg = erro_match.group(1).strip() if erro_match else "Login falhou"
                return f"❌ Falha no login para {empresa}."
            soup = BeautifulSoup(login_response.text, 'html.parser')
            form = soup.find('form', {'name': 'form'})
            if not form:
                return f"❌ Formulário não encontrado após login: {empresa}"
            try:
                pIdSessao = form.find('input', {'name': 'pIdSessao'})['value']
                pNoCache = form.find('input', {'name': 'pNoCache'})['value']
                pPessoa = form.find('input', {'name': 'pPessoa'})['value']
            except (TypeError, KeyError) as e:
                return f"❌ Falha ao extrair parâmetros de sessão para {empresa}."
            select = form.find('select', {'name': 'pObrigacao'})
            if not select:
                return f"❌ Nenhuma fatura disponível"
            pObrigacao = None
            for option in select.find_all('option'):
                if data_vencimento in option.text:
                    pObrigacao = option['value']
                    break
            if not pObrigacao:
                return f"❌ Fatura com vencimento em {data_vencimento} não encontrada."
            fatura_data = {
                "pIdSessao": pIdSessao,
                "pNoCache": pNoCache,
                "pPessoa": pPessoa,
                "pObrigacao": pObrigacao
            }
            request_with_retry(sessao, 'post', FATURA_URL, data=fatura_data)
            confirmacao_data = {
                "pIdSessao": pIdSessao,
                "pNoCache": pNoCache,
                "pPessoa": pPessoa,
                "pObrigacao": pObrigacao,
                "pLinkBoleto": ""
            }
            boleto_response = request_with_retry(sessao, 'post', CONFIRMACAO_URL, data=confirmacao_data)
            url_boleto = boleto_response.url
            if not url_boleto or BASE_URL not in url_boleto:
                return f"❌ URL do boleto inválida para {empresa}"
            nome_arquivo = re.sub(r'[<>:"/\\|?*]', '-', nome_destino)
            pdf_path = os.path.join(save_dir, f"{nome_arquivo}.pdf")
            cookies = [{
                'name': c.name,
                'value': c.value,
                'domain': c.domain,
                'path': c.path,
                'httpOnly': False,
                'secure': False,
                'sameSite': 'Lax'
            } for c in sessao.cookies]
            success = gerar_pdf_com_browser_reutilizavel(context, url_boleto, cookies, pdf_path)
            if success and os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 10000:
                return f"✅ {nome_arquivo}.pdf gerado com sucesso"
            else:
                if os.path.exists(pdf_path):
                    os.remove(pdf_path)
                return f"❌ Falha ao gerar PDF para {empresa}"
        except Exception as e:
            logger.error(f"Erro ao processar boleto para {empresa}: {str(e)}", exc_info=True)
            return f"❌ Erro para {empresa}: {str(e)}"

def inicializar_excel_para_atualizacao(excel_path):
    """Inicializa o workbook do Excel para permitir atualizações"""
    try:
        if not os.path.exists(excel_path):
            logger.error(f"Arquivo Excel não encontrado: {excel_path}")
            return None
        workbook = openpyxl.load_workbook(excel_path)
        sheet = workbook.active
        if sheet.max_row < 2:
            logger.error("Arquivo Excel vazio ou apenas com cabeçalho")
            return None
        status_header = f"Status ({datetime.now().strftime('%d/%m/%Y')})"
        if sheet.cell(row=1, column=STATUS_COLUMN+1).value is None:
            sheet.cell(row=1, column=STATUS_COLUMN+1).value = status_header
            sheet.cell(row=1, column=STATUS_COLUMN+1).font = Font(bold=True)
        workbook.save(excel_path)
        return True
    except Exception as e:
        logger.error(f"Erro ao inicializar Excel: {str(e)}", exc_info=True)
        return None

def atualizar_status_no_excel(excel_path, resultados):
    """Atualiza todos os status de processamento no arquivo Excel de uma só vez"""
    try:
        workbook = openpyxl.load_workbook(excel_path)
        sheet = workbook.active
        for resultado in resultados:
            idx = resultado['index']
            status = resultado['status']
            row = idx + 2
            cell = sheet.cell(row=row, column=STATUS_COLUMN+1)
            
            # Determinar se é sucesso ou falha
            if status.startswith("✅"):
                cell.value = "Sucesso"
                cell.fill = PatternFill(start_color=COR_SUCESSO, end_color=COR_SUCESSO, fill_type="solid")
            else:
                # Extrair o motivo da falha (remover o prefixo "❌" e a parte inicial até ":")
                motivo = status.split(":", 1)[1].strip() if ":" in status else status.replace("❌", "").strip()
                cell.value = f"Falha: {motivo}"
                cell.fill = PatternFill(start_color=COR_ERRO, end_color=COR_ERRO, fill_type="solid")
        
        workbook.save(excel_path)
        return True
    except Exception as e:
        logger.error(f"Erro ao atualizar Excel: {str(e)}")
        return False

def processar_lote_empresas(task_queue, session_pool, data_vencimento, save_dir, progress_callback, worker_id=0):
    """Processa empresas de uma fila de tarefas com um único contexto e pool de sessões"""
    global global_processed
    pw_manager = PlaywrightManager()
    try:
        pw_manager.initialize()
        context = pw_manager.create_context()
        logger.info(f"Worker {worker_id}: Navegador inicializado")
        results = []
        while True:
            try:
                idx, row = task_queue.get_nowait()
                empresa = str(row.get("empresa", "")).strip()
                senha = str(row.get("senha", "")).strip()
                nome = str(row.get("nome", "")).strip().replace("/", "-")
                logger.info(f"Worker {worker_id}: Processando {empresa} ({nome})")
                status = processar_boleto(pw_manager, context, session_pool, empresa, senha, data_vencimento, nome, save_dir)
                logger.info(f"Worker {worker_id}: {status}")
                results.append({
                    'index': idx,
                    'empresa': empresa, 
                    'nome': nome, 
                    'status': status
                })
                task_queue.task_done()
                # Atualizar progresso global de forma thread-safe
                with progress_lock:
                    global_processed += 1
                    percentage = int((global_processed / global_total) * 100)
                    if progress_callback:
                        try:
                            # Chamar o callback com o progresso global
                            progress_callback(percentage, global_processed, global_total)
                        except Exception as e:
                            logger.error(f"Erro ao chamar progress_callback: {str(e)}")
                time.sleep(0.1)  # Reduzido para evitar atrasos desnecessários
            except queue.Empty:
                logger.info(f"Worker {worker_id}: Fila vazia, encerrando")
                break
        context.close()
        return results
    finally:
        pw_manager.close()
        logger.info(f"Worker {worker_id}: Navegador fechado")

def carregar_dados(excel_path):
    """Carrega e valida dados da planilha Excel"""
    try:
        if not os.path.exists(excel_path):
            logger.error(f"Arquivo Excel não encontrado: {excel_path}")
            return None
        df = pd.read_excel(excel_path)
        colunas_requeridas = ["empresa", "senha", "nome"]
        colunas_faltantes = [col for col in colunas_requeridas if col not in df.columns]
        if colunas_faltantes:
            logger.error(f"Colunas faltantes no Excel: {', '.join(colunas_faltantes)}")
            return None
        df = df.dropna(subset=["empresa", "senha"])
        for col in colunas_requeridas:
            if col in df.columns:
                df[col] = df[col].astype(str).str.strip()
        return df
    except Exception as e:
        logger.error(f"Erro ao carregar dados do Excel: {str(e)}", exc_info=True)
        return None

def criar_backup_excel(excel_path):
    """Cria uma cópia de backup do arquivo Excel original"""
    try:
        if not os.path.exists(excel_path):
            return False
        backup_path = f"{excel_path}.bak"
        import shutil
        shutil.copy2(excel_path, backup_path)
        logger.info(f"Backup do Excel criado: {backup_path}")
        return True
    except Exception as e:
        logger.error(f"Erro ao criar backup do Excel: {str(e)}")
        return False

def run_extraction(excel_path, save_dir, due_day, month_number, progress_callback=None):
    """Função principal para extração de boletos, chamada pela interface"""
    global global_processed, global_total
    logger.info("🔄 Iniciando o processamento de boletos Hapvida")
    
    if criar_backup_excel(excel_path):
        logger.info("✅ Backup do arquivo Excel criado com sucesso")
    else:
        logger.warning("⚠️ Não foi possível criar backup do arquivo Excel")
    
    if inicializar_excel_para_atualizacao(excel_path) is None:
        logger.error("❌ Não foi possível preparar o arquivo Excel para atualizações")
        raise Exception("Falha ao preparar o arquivo Excel")
    
    data_vencimento = f"{due_day}/{month_number}/2025"
    if not validar_data_vencimento(data_vencimento):
        logger.error(f"❌ Data de vencimento inválida: {data_vencimento}")
        raise Exception(f"Data de vencimento inválida: {data_vencimento}")
        
    df = carregar_dados(excel_path)
    if df is None or df.empty:
        logger.error("❌ Não foi possível carregar os dados das empresas")
        raise Exception("Falha ao carregar dados das empresas")
        
    logger.info(f"🔍 Processando {len(df)} empresas para boletos com vencimento em {data_vencimento}")
    
    # Validar progress_callback
    if progress_callback and not callable(progress_callback):
        logger.warning("progress_callback não é uma função válida, progresso não será reportado")
        progress_callback = None
    
    session_pool = SessionPool(pool_size=MAX_WORKERS)
    try:
        task_queue = queue.Queue()
        for idx, row in df.iterrows():
            task_queue.put((idx, row))
        
        # Inicializar variáveis globais
        global_processed = 0
        global_total = len(df)
        
        all_results = []
        with ThreadPoolExecutor(max_workers=MAX_WORKERS) as executor:
            futures = {
                executor.submit(
                    processar_lote_empresas, task_queue, session_pool, data_vencimento, save_dir, progress_callback, i
                ): i 
                for i in range(MAX_WORKERS)
            }
            for future in as_completed(futures):
                try:
                    worker_id = futures[future]
                    worker_results = future.result()
                    all_results.extend(worker_results)
                except Exception as e:
                    worker_id = futures[future]
                    logger.error(f"Worker {worker_id} falhou: {str(e)}", exc_info=True)
        
        if atualizar_status_no_excel(excel_path, all_results):
            logger.info("✅ Excel atualizado com sucesso")
        else:
            logger.error("❌ Falha ao atualizar o Excel")
            raise Exception("Falha ao atualizar o Excel")
        
        sucessos = sum(1 for r in all_results if r['status'].startswith("✅"))
        logger.info(f"\n✅ Relatório Final: {sucessos} de {len(all_results)} boletos gerados com sucesso ({sucessos/len(all_results)*100:.1f}%)")
        logger.info(f"📊 O arquivo Excel foi atualizado com os resultados: {excel_path}")
        logger.info(f"📂 PDFs salvos em: {save_dir}")
    
    finally:
        session_pool.close_all()

if __name__ == "__main__":
    # Para testes locais
    run_extraction(
        excel_path=r"C:\Users\chrys\Downloads\SINTICLEPEMP.xlsx",
        save_dir=r"C:\Users\chrys\OneDrive\Desktop\Nova pasta (3)",
        due_day="20",
        month_number="04",
        progress_callback=lambda p, proc, tot: print(f"Progresso: {p}% ({proc}/{tot})")
    )